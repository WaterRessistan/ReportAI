import streamlit as st
import io
from PIL import Image
from docx import Document
from docx.shared import Inches
# IMPORTACIÓN CORRECTA PARA 2026:
from google import genai
from google.genai import types

# --- CONFIGURACIÓN INICIAL ---
st.set_page_config(page_title="ReportAI", layout="wide", page_icon="🛡️")

if "gemini_key" not in st.session_state:
    st.session_state.gemini_key = ""

def analizar_con_gemini(descripcion, normativa, imagenes_archivos):
    # Usamos el nuevo cliente de Google GenAI
    client = genai.Client(api_key=st.session_state.gemini_key.strip())
    
    # Preparamos las imágenes y el texto para el nuevo formato
    # Importante: El prompt principal va como primer elemento
    contenido_multimodal = [f"Normativa de referencia: {normativa}. Hallazgo descrito: {descripcion}"]
    
    for f in imagenes_archivos:
        bytes_data = f.getvalue()
        contenido_multimodal.append(
            types.Part.from_bytes(data=bytes_data, mime_type="image/jpeg")
        )

    # El nuevo método de generación
    response = client.models.generate_content(
        model='gemini-2.5-flash', # Google ya resuelve la ruta internamente
        contents=contenido_multimodal,
        config=types.GenerateContentConfig(
            system_instruction="""Actúa como un auditor de seguridad informática senior. 
            Analiza las evidencias. Si falta info responde [INCOMPLETO]. 
            Si está todo, empieza con [COMPLETO] y redacta el informe profesional."""
        )
    )
    
    return response.text

def exportar_a_word(texto_ia, imagenes):
    doc = Document()
    doc.add_heading('INFORME DE AUDITORÍA DE SEGURIDAD', 0)
    
    # Estilo básico para que el informe no sea un bloque de texto plano
    doc.add_heading('Análisis de la IA', level=1)
    doc.add_paragraph(texto_ia)
    
    if imagenes:
        doc.add_heading('Evidencias Visuales Adjuntas', level=1)
        for i, img in enumerate(imagenes):
            img.seek(0)
            doc.add_picture(img, width=Inches(4))
            doc.add_paragraph(f"Anexo {i+1}: Captura de evidencia.")
            
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- INTERFAZ DE USUARIO ---

with st.sidebar:
    st.title("🔑 Acceso")
    key = st.text_input("Introduce tu Gemini API Key:", type="password")
    if key:
        st.session_state.gemini_key = key
    st.markdown("[Consigue tu clave gratis aquí](https://aistudio.google.com/)")

st.title("🛡️ ReportAI v1.0")
st.info("Analizador de vulnerabilidades mediante visión artificial.")

col1, col2 = st.columns([1, 1])

with col1:
    desc = st.text_area("Descripción del hallazgo:", 
                        placeholder="Ej: Acceso SSH exitoso con credenciales por defecto...", 
                        height=150)
    norma = st.selectbox("Marco Normativo:", ["ISO 27001", "NIST SP 800-53", "ENS España", "GDPR"])

with col2:
    fotos = st.file_uploader("Subir capturas de pantalla:", accept_multiple_files=True, type=['png', 'jpg', 'jpeg'])

if st.button("🚀 Iniciar Auditoría"):
    if not st.session_state.gemini_key:
        st.error("Falta la API Key.")
    elif not desc or not fotos:
        st.warning("Por favor, aporta una descripción y al menos una imagen.")
    else:
        with st.spinner("La IA está procesando las evidencias..."):
            try:
                resultado = analizar_con_gemini(desc, norma, fotos)
                
                if "[INCOMPLETO]" in resultado:
                    st.warning("⚠️ Información insuficiente según la IA:")
                    st.write(resultado.replace("[INCOMPLETO]", ""))
                else:
                    st.session_state.resultado_final = resultado.replace("[COMPLETO]", "")
                    st.success("✅ Informe generado con éxito.")
            except Exception as e:
                st.error(f"Error técnico: {e}")

if "resultado_final" in st.session_state:
    st.divider()
    st.markdown(st.session_state.resultado_final)
    
    btn_word = exportar_a_word(st.session_state.resultado_final, fotos)
    st.download_button("📥 Descargar Reporte (.docx)", 
                       data=btn_word, 
                       file_name="Reporte_Auditoria.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

st.markdown("---")
st.markdown("Made by WaterRessistan")